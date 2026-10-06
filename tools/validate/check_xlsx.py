"""Checks a plain .xlsx export against formats.md §6 and reads it with openpyxl.

1. The package: exactly the seven parts of formats.md §6.1 (no docProps), every entry deflated, and the
   same fixed local time on every entry, 1980-01-01 00:00, with no extra field that carries a time
   (formats-review P9: a time set through the default zone reveals the phone's UTC offset).
2. The sheet: every cell a shared string (t="s"), no formula, no inline string, no hyperlink and no
   external relationship; the counts of sharedStrings.xml right.
3. openpyxl 3.1.5: one sheet of 11 columns, frozen header row, print titles, filter, bold header, text
   format "@" on every data cell, no formula cell; the column widths are printed for information.

Ported from prototypes/formats/validate/check_xlsx.py. openpyxl does not decode the _xHHHH_ escapes of
characters XML forbids and deletes "x005F_" (formats-review P17.5), so cell values are compared in
compare_cells.py, where LibreOffice's reading is the reference.
Usage: check_xlsx.py [--show-values] FILE.xlsx
"""
import argparse
import os
import re
import struct
import sys
import zipfile

from common import Report

PARTS = {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml", "xl/_rels/workbook.xml.rels", "xl/styles.xml",
         "xl/worksheets/sheet1.xml", "xl/sharedStrings.xml"}
FIXED_TIME = (1980, 1, 1, 0, 0, 0)
WIDTHS = [24, 26, 22, 30, 24, 40, 30, 10, 16, 16, 16]  # formats.md §6.2
TIME_EXTRAS = {0x5455: "extended timestamp", 0x000A: "NTFS times", 0x5855: "Info-ZIP Unix times"}


def extra_ids(extra):
    ids, i = [], 0
    while i + 4 <= len(extra):
        hid, size = struct.unpack("<HH", extra[i:i + 4])
        ids.append(hid)
        i += 4 + size
    return ids


def check(path, report):
    """Checks one workbook; returns openpyxl's cell values (rows of strings, None for an empty cell), or None."""
    name = os.path.basename(path)
    try:
        z = zipfile.ZipFile(path)
    except zipfile.BadZipFile as e:
        report.fail(f"{name}: a zip package", [str(e)])
        return None
    infos = z.infolist()
    names = {i.filename for i in infos}
    report.check(names == PARTS and len(infos) == len(PARTS), f"{name}: exactly the seven parts of formats.md §6.1",
                 [f"missing {sorted(PARTS - names)}" if PARTS - names else "", f"extra {sorted(names - PARTS)}" if names - PARTS else ""])
    report.check(all(i.compress_type == zipfile.ZIP_DEFLATED for i in infos), f"{name}: every entry deflated")
    times = sorted({i.date_time for i in infos})
    extras = sorted({TIME_EXTRAS[h] for i in infos for h in extra_ids(i.extra) if h in TIME_EXTRAS})
    report.check(times == [FIXED_TIME] and not extras,
                 f"{name}: every entry has the fixed local time 1980-01-01 00:00 and no time field (formats-review P9)",
                 [f"entry times {times}"] + ([f"extra fields with times: {extras}"] if extras else []))
    sheet = z.read("xl/worksheets/sheet1.xml").decode("utf-8") if "xl/worksheets/sheet1.xml" in names else ""
    cells = re.findall(r"<c\b([^>]*?)(/?)>", sheet)
    not_shared = [a for a, _ in cells if 't="s"' not in a]
    report.check(cells and not not_shared, f"{name}: every cell is a shared string (formats.md §6.2)  [{len(cells)} cells]",
                 [f"{len(not_shared)} cells without t=\"s\", the first: <c{not_shared[0]}>" if not_shared else "no cells"])
    banned = [w for w, pat in (("formula", r"<f[\s>/]"), ("inline string", r"<is>"), ("hyperlink", r"<hyperlinks?[\s>]"))
              if re.search(pat, sheet)]
    external = [n for n in names if n.endswith(".rels") and "TargetMode=\"External\"" in z.read(n).decode("utf-8")]
    report.check(not banned and not external, f"{name}: no formula, inline string, hyperlink or external relationship",
                 banned + [f"external target in {n}" for n in external])
    if "xl/sharedStrings.xml" in names:
        sst = z.read("xl/sharedStrings.xml").decode("utf-8")
        m = re.search(r"<sst\b[^>]*>", sst)
        head = m.group(0) if m else ""
        count = re.search(r'\bcount="(\d+)"', head)
        unique = re.search(r'\buniqueCount="(\d+)"', head)
        n_si = len(re.findall(r"<si>", sst))
        report.check(bool(count and unique) and int(count.group(1)) == len(cells) and int(unique.group(1)) == n_si,
                     f"{name}: sharedStrings count and uniqueCount are right",
                     [f"count={count and count.group(1)} for {len(cells)} cells, uniqueCount={unique and unique.group(1)} for {n_si} strings"])
    try:
        import openpyxl
    except ImportError:
        report.skip("openpyxl is missing: run tools/validate/setup.sh")
        return None
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    report.check(len(wb.worksheets) == 1 and ws.max_column == 11, f"{name}: one sheet of 11 columns",
                 [f"{len(wb.worksheets)} sheets, {ws.max_column} columns"])
    setup = []
    if ws.freeze_panes != "A2":
        setup.append(f"frozen pane {ws.freeze_panes}, expected A2")
    if (ws.print_title_rows or "").replace("$", "") != "1:1":
        setup.append(f"print titles {ws.print_title_rows}, expected 1:1")
    if not (ws.auto_filter.ref or "").startswith("A1:"):
        setup.append(f"filter {ws.auto_filter.ref}, expected A1:…")
    if not all(c.font.b for c in ws[1]):
        setup.append("the header row is not bold")
    report.check(not setup, f"{name}: frozen header, print titles, filter, bold header", setup)
    values, not_text, formulas = [], 0, 0
    for r, row in enumerate(ws.iter_rows(), start=1):
        vals = []
        for c in row:
            if c.data_type == "f":
                formulas += 1
            if r > 1 and c.value is not None and c.number_format != "@":  # empty cells are not written
                not_text += 1
            vals.append(c.value)
        values.append(vals)
    report.check(not_text == 0 and formulas == 0, f"{name}: every written data cell has the text format \"@\" and none is a formula",
                 [f"{not_text} cells without \"@\", {formulas} formula cells"])
    widths = [ws.column_dimensions[chr(ord("A") + i)].width for i in range(11)]
    report.info(f"{name}: sheet {ws.title!r}, {ws.max_row - 1} rows under the header; column widths "
                + ("as formats.md §6.2" if [round(w or 0) for w in widths] == WIDTHS else f"{widths}, formats.md §6.2 says {WIDTHS}"))
    return values


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("--show-values", action="store_true")
    a = ap.parse_args()
    r = Report(a.show_values)
    check(a.file, r)
    sys.exit(1 if r.failed else 0)


if __name__ == "__main__":
    main()
